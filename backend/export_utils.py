"""
CSV / JSON / Excel export of analysis results and model comparisons.
Input is the JSON the API returns from /api/predict and /api/compare.
"""

from __future__ import annotations

import csv
import io
import json
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

HEADERS = {
    "en": ["Timestamp", "Image", "Predicted class", "Confidence", "Model", "Explanation", "Mode", "Inference (ms)",
           "Top-2 class", "Top-2 conf.", "Top-3 class", "Top-3 conf.", "Severity"],
    "ru": ["Время", "Изображение", "Предсказанный класс", "Уверенность", "Модель", "Объяснение", "Режим", "Инференс (мс)",
           "Класс №2", "Уверенность №2", "Класс №3", "Уверенность №3", "Серьёзность"],
}
COMPARISON_HEADERS = {
    "en": ["Model", "Predicted class", "Confidence", "Inference (ms)", "Trained"],
    "ru": ["Модель", "Предсказанный класс", "Уверенность", "Инференс (мс)", "Обучена"],
}


def _cls(entry: dict, lang: str) -> str:
    return entry.get("class_ru" if lang == "ru" else "class") or entry.get("class", "")


def _row(result: dict, lang: str) -> list:
    pred = result.get("prediction", {})
    top = result.get("top_predictions", [])
    tr = result.get("treatment") or {}
    row = [
        result.get("timestamp", datetime.now().isoformat(timespec="seconds")),
        result.get("image_name", ""),
        _cls(pred, lang),
        float(pred.get("confidence", 0)),
        result.get("model_label") or result.get("model", ""),
        result.get("explanation", ""),
        result.get("dataset_type", ""),
        result.get("inference_time_ms", ""),
    ]
    for i in (1, 2):
        if i < len(top):
            row += [_cls(top[i], lang), float(top[i].get("confidence", 0))]
        else:
            row += ["", ""]
    row.append(tr.get("severity", ""))
    return row


def results_to_csv(results: list[dict], lang: str = "en") -> str:
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(HEADERS[lang])
    for r in results:
        w.writerow(_row(r, lang))
    return out.getvalue()


def results_to_json(results: list[dict]) -> str:
    slim = []
    for r in results:
        slim.append({k: v for k, v in r.items() if k != "images"})
    return json.dumps({"exported_at": datetime.now().isoformat(timespec="seconds"), "total": len(slim), "results": slim}, ensure_ascii=False, indent=2)


def _style_sheet(ws, headers, rows, percent_cols=()):
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill(start_color="2F6B3A", end_color="2F6B3A", fill_type="solid")
    thin = Side(style="thin", color="D0D7DE")
    border = Border(left=thin, right=thin, top=thin, bottom=thin)
    for c, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.font, cell.fill, cell.border = head_font, head_fill, border
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for r, row in enumerate(rows, 2):
        for c, value in enumerate(row, 1):
            cell = ws.cell(row=r, column=c, value=value)
            cell.border = border
            if c in percent_cols and isinstance(value, (int, float)):
                cell.number_format = "0.00%"
    for c in range(1, len(headers) + 1):
        width = max(len(str(ws.cell(row=r, column=c).value or "")) for r in range(1, len(rows) + 2))
        ws.column_dimensions[get_column_letter(c)].width = min(max(10, width + 2), 48)
    ws.freeze_panes = "A2"


def results_to_excel(results: list[dict], lang: str = "en") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Results" if lang == "en" else "Результаты"
    _style_sheet(ws, HEADERS[lang], [_row(r, lang) for r in results], percent_cols=(4, 10, 12))

    ws2 = wb.create_sheet("Top-5")
    hdr = ["Image", "Rank", "Class", "Confidence"] if lang == "en" else ["Изображение", "Ранг", "Класс", "Уверенность"]
    rows = []
    for r in results:
        for i, t in enumerate(r.get("top_predictions", []), 1):
            rows.append([r.get("image_name", ""), i, _cls(t, lang), float(t.get("confidence", 0))])
    _style_sheet(ws2, hdr, rows, percent_cols=(4,))

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _comparison_rows(data: dict, lang: str) -> list[list]:
    rows = []
    for name, r in data.get("comparisons", {}).items():
        rows.append(
            [
                r.get("label") or name,
                _cls(r, lang),
                float(r.get("confidence", 0)),
                r.get("inference_time_ms", ""),
                ("да" if lang == "ru" else "yes") if r.get("trained", True) else ("нет" if lang == "ru" else "no"),
            ]
        )
    return rows


def comparison_to_csv(data: dict, lang: str = "en") -> str:
    out = io.StringIO()
    w = csv.writer(out, delimiter=";")
    w.writerow(["Image" if lang == "en" else "Изображение", data.get("image_name", "")])
    w.writerow(COMPARISON_HEADERS[lang])
    for row in _comparison_rows(data, lang):
        w.writerow(row)
    return out.getvalue()


def comparison_to_excel(data: dict, lang: str = "en") -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "Comparison" if lang == "en" else "Сравнение"
    _style_sheet(ws, COMPARISON_HEADERS[lang], _comparison_rows(data, lang), percent_cols=(3,))
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()
